"""Tests for the approved status feature (FEAT-002)."""
from __future__ import annotations

from pathlib import Path

import pytest
from typer.testing import CliRunner

from stx.cli import app
from stx.excel import export_document_to_excel, import_document_from_excel
from stx.model import Document, Entry
from stx.stf import parse_stf_text, render_full_stf
from stx.validate import validate_document

runner = CliRunner()


# ---------------------------------------------------------------------------
# Entry model tests
# ---------------------------------------------------------------------------


def test_entry_approved_defaults_to_false():
    e = Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa")
    assert e.approved is False


def test_entry_approved_true_with_translation_returns_approved_status():
    e = Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa", approved=True)
    assert e.status == "Approved"


def test_entry_approved_true_without_translation_returns_untranslated():
    e = Entry(key="CustomLabel.A", label="Hello", translation="", approved=True)
    assert e.status == "Untranslated"


def test_entry_logical_sheet_name_ignores_approved():
    """Approved entries should group as Translated, not Approved."""
    e = Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa", approved=True)
    assert e.logical_sheet_name == "CustomLabel_Translated"


# ---------------------------------------------------------------------------
# Excel round-trip
# ---------------------------------------------------------------------------


def test_excel_round_trip_preserves_approved(tmp_path):
    doc = Document(
        language="Japanese",
        language_code="ja",
        entries=[
            Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa", approved=True),
            Entry(key="CustomLabel.B", label="World", translation="Sekai", approved=False),
            Entry(key="CustomLabel.C", label="Bye"),
        ],
    )
    xlsx = tmp_path / "test.xlsx"
    export_document_to_excel(doc, xlsx)
    loaded = import_document_from_excel(xlsx)

    assert loaded.entries[0].approved is True
    assert loaded.entries[1].approved is False
    assert loaded.entries[2].approved is False


def test_excel_backward_compatible_no_approved_column(tmp_path):
    """Workbooks without an Approved column should import fine (all False)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "CustomLabel_Translated"
    ws.append(["Key", "Label", "Translation"])
    ws.append(["CustomLabel.A", "Hello", "Konnichiwa"])
    ws.append(["CustomLabel.B", "World", "Sekai"])
    xlsx = tmp_path / "old_format.xlsx"
    wb.save(xlsx)

    loaded = import_document_from_excel(xlsx)
    assert all(e.approved is False for e in loaded.entries)


# ---------------------------------------------------------------------------
# STF round-trip
# ---------------------------------------------------------------------------


def test_stf_round_trip_preserves_approved():
    doc = Document(
        language="Japanese",
        language_code="ja",
        entries=[
            Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa", approved=True),
            Entry(key="CustomLabel.B", label="World", translation="Sekai", approved=False),
            Entry(key="CustomLabel.C", label="Bye"),
        ],
    )
    text = render_full_stf(doc)
    # Verify the marker appears in text
    assert "# APPROVED" in text

    parsed = parse_stf_text(text)
    assert parsed.entries[0].approved is True
    assert parsed.entries[1].approved is False
    assert parsed.entries[2].approved is False


def test_stf_approved_marker_only_for_translated():
    """Untranslated entries should never get the # APPROVED marker."""
    doc = Document(
        language="Japanese",
        language_code="ja",
        entries=[
            Entry(key="CustomLabel.A", label="Hello", translation="", approved=True),
        ],
    )
    text = render_full_stf(doc)
    assert "# APPROVED" not in text


# ---------------------------------------------------------------------------
# Validator
# ---------------------------------------------------------------------------


def test_validator_skips_approved_entries():
    """Approved entries should not produce per-entry issues."""
    doc = Document(
        language="Japanese",
        language_code="ja",
        entries=[
            # This would normally trigger a length_limit error (>80 chars for CustomField)
            Entry(
                key="CustomField.Long.FieldLabel",
                label="A" * 100,
                translation="B" * 100,
                approved=True,
            ),
        ],
    )
    report = validate_document(doc)
    # Should only have the info about skipped entries, not a length_limit error
    categories = {i.category for i in report.issues}
    assert "length_limit" not in categories
    assert "approved_skipped" in categories
    # Check the info message
    info_issues = [i for i in report.issues if i.category == "approved_skipped"]
    assert len(info_issues) == 1
    assert "1" in info_issues[0].message


def test_validator_still_catches_duplicate_keys_for_approved():
    """Duplicate key checks are document-level and should not be skipped."""
    doc = Document(
        language="Japanese",
        language_code="ja",
        entries=[
            Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa", approved=True),
            Entry(key="CustomLabel.A", label="World", translation="Sekai", approved=True),
        ],
    )
    report = validate_document(doc)
    categories = {i.category for i in report.issues}
    assert "duplicate_key" in categories


# ---------------------------------------------------------------------------
# CLI approve / unapprove
# ---------------------------------------------------------------------------


def _write_sample_stf(tmp_path: Path) -> Path:
    doc = Document(
        language="Japanese",
        language_code="ja",
        entries=[
            Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa"),
            Entry(key="CustomLabel.B", label="World", translation="Sekai"),
            Entry(key="CustomLabel.C", label="Bye"),
        ],
    )
    path = tmp_path / "test.stf"
    path.write_text(render_full_stf(doc), encoding="utf-8")
    return path


def test_cli_approve_by_keys(tmp_path):
    stf = _write_sample_stf(tmp_path)
    result = runner.invoke(app, ["approve", str(stf), "--keys", "CustomLabel.A,CustomLabel.B"])
    assert result.exit_code == 0
    assert "Approved 2" in result.output

    # Verify the file was updated
    parsed = parse_stf_text(stf.read_text(encoding="utf-8"))
    assert parsed.entries[0].approved is True
    assert parsed.entries[1].approved is True
    assert parsed.entries[2].approved is False


def test_cli_approve_all_translated(tmp_path):
    stf = _write_sample_stf(tmp_path)
    result = runner.invoke(app, ["approve", str(stf), "--all-translated"])
    assert result.exit_code == 0
    assert "Approved 2" in result.output

    parsed = parse_stf_text(stf.read_text(encoding="utf-8"))
    assert parsed.entries[0].approved is True
    assert parsed.entries[1].approved is True
    assert parsed.entries[2].approved is False


def test_cli_unapprove_by_keys(tmp_path):
    # First approve, then unapprove
    doc = Document(
        language="Japanese",
        language_code="ja",
        entries=[
            Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa", approved=True),
            Entry(key="CustomLabel.B", label="World", translation="Sekai", approved=True),
        ],
    )
    stf = tmp_path / "test.stf"
    stf.write_text(render_full_stf(doc), encoding="utf-8")

    result = runner.invoke(app, ["unapprove", str(stf), "--keys", "CustomLabel.A"])
    assert result.exit_code == 0
    assert "Unapproved 1" in result.output

    parsed = parse_stf_text(stf.read_text(encoding="utf-8"))
    assert parsed.entries[0].approved is False
    assert parsed.entries[1].approved is True


def test_cli_unapprove_all(tmp_path):
    doc = Document(
        language="Japanese",
        language_code="ja",
        entries=[
            Entry(key="CustomLabel.A", label="Hello", translation="Konnichiwa", approved=True),
            Entry(key="CustomLabel.B", label="World", translation="Sekai", approved=True),
        ],
    )
    stf = tmp_path / "test.stf"
    stf.write_text(render_full_stf(doc), encoding="utf-8")

    result = runner.invoke(app, ["unapprove", str(stf), "--all"])
    assert result.exit_code == 0
    assert "Unapproved 2" in result.output

    parsed = parse_stf_text(stf.read_text(encoding="utf-8"))
    assert all(not e.approved for e in parsed.entries)


def test_cli_approve_no_flags(tmp_path):
    """Should fail if neither --keys nor --all-translated is provided."""
    stf = _write_sample_stf(tmp_path)
    result = runner.invoke(app, ["approve", str(stf)])
    assert result.exit_code == 2


def test_cli_unapprove_no_flags(tmp_path):
    """Should fail if neither --keys nor --all is provided."""
    stf = _write_sample_stf(tmp_path)
    result = runner.invoke(app, ["unapprove", str(stf)])
    assert result.exit_code == 2
