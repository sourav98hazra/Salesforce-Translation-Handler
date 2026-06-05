"""Tests for the import_translations module."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from stx.import_translations import ImportedTranslations, parse_translation_file
from stx.model import Document, Entry
from stx.translate import translate_document
from stx.translate.base import Translator


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _write_workbook(path: Path, sheets: dict[str, list[list]]) -> Path:
    """Helper to write a multi-sheet .xlsx for testing."""
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Test: ImportedTranslations class
# ---------------------------------------------------------------------------


class TestImportedTranslations:
    def test_empty(self) -> None:
        it = ImportedTranslations()
        assert it.count == 0
        assert it.get("anything") is None

    def test_with_data(self) -> None:
        it = ImportedTranslations(translations={"Hello": "Bonjour", "World": "Monde"})
        assert it.count == 2
        assert it.get("Hello") == "Bonjour"
        assert it.get("World") == "Monde"
        assert it.get("Missing") is None


# ---------------------------------------------------------------------------
# Test: Standard format (Key / Label / Translation columns)
# ---------------------------------------------------------------------------


class TestParseStandardFormat:
    def test_standard_columns(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "standard.xlsx", {
            "Sheet1": [
                ["Key", "Label", "Translation", "Approved"],
                ["Key1", "Hello", "Hola", ""],
                ["Key2", "World", "Mundo", ""],
                ["Key3", "Empty", "", ""],  # no translation -- should be skipped
            ]
        })
        result = parse_translation_file(path)
        assert result.count == 2
        assert result.get("Hello") == "Hola"
        assert result.get("World") == "Mundo"
        assert result.get("Empty") is None

    def test_multiple_sheets(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "multi.xlsx", {
            "Sheet1": [
                ["Key", "Label", "Translation"],
                ["K1", "Hello", "Bonjour"],
            ],
            "Sheet2": [
                ["Key", "Label", "Translation"],
                ["K2", "World", "Monde"],
                # Duplicate from Sheet1 -- first occurrence wins
                ["K3", "Hello", "Salut"],
            ],
        })
        result = parse_translation_file(path)
        assert result.count == 2
        assert result.get("Hello") == "Bonjour"  # first sheet wins
        assert result.get("World") == "Monde"


# ---------------------------------------------------------------------------
# Test: Glossary-style (Source / Translation or just 2 columns)
# ---------------------------------------------------------------------------


class TestParseGlossaryStyle:
    def test_source_translation_columns(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "glossary.xlsx", {
            "Sheet1": [
                ["Source", "Translation"],
                ["Good morning", "Buenos dias"],
                ["Good night", "Buenas noches"],
            ]
        })
        result = parse_translation_file(path)
        assert result.count == 2
        assert result.get("Good morning") == "Buenos dias"
        assert result.get("Good night") == "Buenas noches"

    def test_source_target_columns(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "src_tgt.xlsx", {
            "Sheet1": [
                ["Source", "Target"],
                ["Cat", "Gato"],
                ["Dog", "Perro"],
            ]
        })
        result = parse_translation_file(path)
        assert result.count == 2
        assert result.get("Cat") == "Gato"

    def test_two_column_fallback(self, tmp_path: Path) -> None:
        """When there are exactly 2 non-empty header columns, use them."""
        path = _write_workbook(tmp_path / "two_col.xlsx", {
            "Sheet1": [
                ["English", "Spanish"],
                ["Apple", "Manzana"],
                ["Banana", "Platano"],
            ]
        })
        result = parse_translation_file(path)
        assert result.count == 2
        assert result.get("Apple") == "Manzana"
        assert result.get("Banana") == "Platano"


# ---------------------------------------------------------------------------
# Test: Custom column specification
# ---------------------------------------------------------------------------


class TestCustomColumns:
    def test_custom_source_and_translation_cols(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "custom.xlsx", {
            "Sheet1": [
                ["ID", "OriginalText", "TranslatedText", "Notes"],
                ["1", "Hello", "Hallo", "German"],
                ["2", "World", "Welt", "German"],
            ]
        })
        result = parse_translation_file(
            path, source_col="OriginalText", translation_col="TranslatedText"
        )
        assert result.count == 2
        assert result.get("Hello") == "Hallo"
        assert result.get("World") == "Welt"

    def test_custom_cols_case_insensitive(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "case.xlsx", {
            "Sheet1": [
                ["ENGLISH", "JAPANESE"],
                ["Tree", "Ki"],
            ]
        })
        result = parse_translation_file(
            path, source_col="english", translation_col="japanese"
        )
        assert result.count == 1
        assert result.get("Tree") == "Ki"


# ---------------------------------------------------------------------------
# Test: Empty / missing file handling
# ---------------------------------------------------------------------------


class TestEdgeCases:
    def test_missing_file(self, tmp_path: Path) -> None:
        result = parse_translation_file(tmp_path / "nonexistent.xlsx")
        assert result.count == 0

    def test_empty_sheet(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "empty.xlsx", {
            "Sheet1": []
        })
        result = parse_translation_file(path)
        assert result.count == 0

    def test_no_matching_columns(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "bad_cols.xlsx", {
            "Sheet1": [
                ["Col1", "Col2", "Col3", "Col4"],
                ["a", "b", "c", "d"],
            ]
        })
        result = parse_translation_file(path)
        assert result.count == 0

    def test_none_cells_ignored(self, tmp_path: Path) -> None:
        path = _write_workbook(tmp_path / "nones.xlsx", {
            "Sheet1": [
                ["Label", "Translation"],
                [None, "ShouldBeSkipped"],
                ["Source", None],
                ["Valid", "ValidTranslation"],
            ]
        })
        result = parse_translation_file(path)
        assert result.count == 1
        assert result.get("Valid") == "ValidTranslation"


# ---------------------------------------------------------------------------
# Test: Integration with runner
# ---------------------------------------------------------------------------


class CountingTranslator(Translator):
    """Test double that records every call."""

    def __init__(self) -> None:
        self.calls: list[tuple[str, str, str]] = []

    def translate(self, text: str, source_lang: str, target_lang: str) -> str:
        self.calls.append((text, source_lang, target_lang))
        return f"<{target_lang}>{text}</{target_lang}>"


class TestRunnerIntegration:
    def test_imported_translations_have_highest_priority(self, tmp_path: Path) -> None:
        """Imported translations should be used before TM and before network."""
        from stx.memory import TranslationMemory

        # Set up TM with a different translation for "Hello"
        tm_path = tmp_path / "tm.sqlite"
        tm = TranslationMemory(path=tm_path)
        tm.put("Hello", "en", "ja", "TM-Hello")

        doc = Document(entries=[
            Entry(key="K1", label="Hello"),
            Entry(key="K2", label="World"),
            Entry(key="K3", label="NotImported"),
        ])

        imported = {"Hello": "Imported-Hello", "World": "Imported-World"}
        translator = CountingTranslator()

        result = translate_document(
            doc, translator,
            source_lang="en", target_lang="ja",
            memory=tm,
            workers=1, rate_limit_per_second=None,
            prevent_system_sleep=False,
            imported_translations=imported,
        )

        # "Hello" and "World" should come from imported, not TM or network
        assert result.imported_reuse_count == 2
        # Only "NotImported" should go through the translator
        assert len(translator.calls) == 1
        assert translator.calls[0][0] == "NotImported"

        # Verify the translations are correct
        assert doc.entries[0].translation == "Imported-Hello"
        assert doc.entries[1].translation == "Imported-World"
        assert doc.entries[2].translation == "<ja>NotImported</ja>"

    def test_imported_overrides_already_translated_without_retranslate(self) -> None:
        """Imported translations override already-translated rows even without retranslate_existing=True."""
        doc = Document(entries=[
            Entry(key="K1", label="Hello", translation="OldHello"),
            Entry(key="K2", label="World", translation="OldWorld"),
            Entry(key="K3", label="Foo", translation="OldFoo"),
        ])

        imported = {"Hello": "Imported-Hello", "World": "Imported-World"}
        translator = CountingTranslator()

        result = translate_document(
            doc, translator,
            source_lang="en", target_lang="ja",
            workers=1, rate_limit_per_second=None,
            prevent_system_sleep=False,
            imported_translations=imported,
            retranslate_existing=False,
        )

        # K1 and K2 have imported matches so they bypass the skip gate
        assert doc.entries[0].translation == "Imported-Hello"
        assert doc.entries[1].translation == "Imported-World"
        # K3 has no imported match and retranslate_existing=False, so it stays
        assert doc.entries[2].translation == "OldFoo"
        assert result.imported_reuse_count == 2
        assert result.skipped_count == 1
        # Imported rows are not counted as retranslated
        assert result.retranslated_count == 0

    def test_imported_none_means_no_effect(self) -> None:
        """When imported_translations is None, normal flow is used."""
        doc = Document(entries=[
            Entry(key="K1", label="Hello"),
        ])
        translator = CountingTranslator()

        result = translate_document(
            doc, translator,
            source_lang="en", target_lang="ja",
            workers=1, rate_limit_per_second=None,
            prevent_system_sleep=False,
            imported_translations=None,
        )

        assert result.imported_reuse_count == 0
        assert len(translator.calls) == 1

    def test_imported_dedup_propagation(self) -> None:
        """Imported translations populate the dedup cache for repeated labels."""
        doc = Document(entries=[
            Entry(key="K1", label="Hello"),
            Entry(key="K2", label="Hello"),
            Entry(key="K3", label="Hello"),
        ])
        imported = {"Hello": "Imported-Hello"}
        translator = CountingTranslator()

        result = translate_document(
            doc, translator,
            source_lang="en", target_lang="ja",
            workers=1, rate_limit_per_second=None,
            prevent_system_sleep=False,
            imported_translations=imported,
        )

        # All occurrences are served from imported (first is imported,
        # subsequent find it in the dedup cache populated by the first).
        assert result.imported_reuse_count >= 1
        assert result.translated_count == 3
        assert len(translator.calls) == 0
        # All should have the imported translation
        for entry in doc.entries:
            assert entry.translation == "Imported-Hello"
